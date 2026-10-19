"""
Report Generation Service for AI-Based Training Centre Compliance & Monitoring System.
Generates genuine PDF (via ReportLab) and Excel (via OpenPyXL) reports.
"""

import io
import os
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from reportlab.lib.pagesizes import letter, A4
from reportlab.lib import colors
from reportlab.lib.units import inch
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

def generate_attendance_excel(centre_data, session_data, records_data):
    """Generate professional Excel workbook for Attendance Report."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Attendance Compliance"

    # Styling definitions
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid") # Navy
    sub_fill = PatternFill(start_color="3B82F6", end_color="3B82F6", fill_type="solid")
    alt_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    warn_fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    title_font = Font(name="Calibri", size=16, bold=True, color="1E3A8A")
    bold_font = Font(name="Calibri", size=11, bold=True)
    regular_font = Font(name="Calibri", size=11)
    
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    # Title & Metadata
    ws['A1'] = "GOVERNMENT OF INDIA - SKILLING SCHEME MONITORING SYSTEM"
    ws['A1'].font = Font(name="Calibri", size=14, bold=True, color="475569")
    ws['A2'] = f"ATTENDANCE COMPLIANCE & AI CROSS-CHECK REPORT"
    ws['A2'].font = title_font

    ws['A4'] = "Training Centre:"
    ws['B4'] = centre_data.get('centre_name', 'N/A')
    ws['A4'].font = bold_font
    ws['B4'].font = regular_font

    ws['A5'] = "Centre ID & Location:"
    ws['B5'] = f"{centre_data.get('centre_id', '')} - {centre_data.get('location', '')}, {centre_data.get('state', '')}"
    ws['A5'].font = bold_font
    ws['B5'].font = regular_font

    ws['D4'] = "Date of Report:"
    ws['E4'] = datetime.now().strftime("%Y-%m-%d %H:%M")
    ws['D4'].font = bold_font
    ws['E4'].font = regular_font

    ws['D5'] = "Verification Status:"
    ws['E5'] = session_data.get('discrepancy_status', 'NORMAL')
    ws['D5'].font = bold_font
    ws['E5'].font = Font(name="Calibri", size=11, bold=True, color="B91C1C" if session_data.get('discrepancy_status') != 'NORMAL' else "15803D")

    # Summary Metrics Box
    ws['A7'] = "METRIC"
    ws['B7'] = "MANUAL LOG"
    ws['C7'] = "AI CAMERA HEADCOUNT"
    ws['D7'] = "DIFFERENCE"
    ws['E7'] = "DISCREPANCY %"

    for col in ['A7', 'B7', 'C7', 'D7', 'E7']:
        ws[col].fill = sub_fill
        ws[col].font = header_font
        ws[col].alignment = Alignment(horizontal="center")

    ws['A8'] = "Trainee Headcount"
    ws['B8'] = session_data.get('manual_present_count', 0)
    ws['C8'] = session_data.get('ai_detected_count', 0)
    ws['D8'] = session_data.get('difference_count', 0)
    ws['E8'] = f"{session_data.get('discrepancy_percentage', 0.0):.1f}%"

    for col in ['A8', 'B8', 'C8', 'D8', 'E8']:
        ws[col].font = bold_font
        ws[col].alignment = Alignment(horizontal="center")
        ws[col].border = thin_border
        if col in ['D8', 'E8'] and session_data.get('difference_count', 0) > 0:
            ws[col].fill = warn_fill

    # Trainees List Table
    ws['A10'] = "S.No"
    ws['B10'] = "Trainee ID"
    ws['C10'] = "Trainee Name"
    ws['D10'] = "Manual Status"
    ws['E10'] = "Time Marked"
    ws['F10'] = "Remarks"

    for col in ['A10', 'B10', 'C10', 'D10', 'E10', 'F10']:
        ws[col].fill = header_fill
        ws[col].font = header_font
        ws[col].alignment = Alignment(horizontal="center" if col in ['A10', 'B10', 'D10', 'E10'] else "left")

    row_idx = 11
    for i, r in enumerate(records_data, 1):
        ws[f'A{row_idx}'] = i
        ws[f'B{row_idx}'] = r.get('trainee_id', '')
        ws[f'C{row_idx}'] = r.get('name', '')
        ws[f'D{row_idx}'] = r.get('status', 'Present')
        ws[f'E{row_idx}'] = r.get('marked_time', '—')
        ws[f'F{row_idx}'] = r.get('remarks', '')

        fill = alt_fill if i % 2 == 0 else PatternFill(fill_type=None)
        for col_letter in ['A', 'B', 'C', 'D', 'E', 'F']:
            cell = ws[f'{col_letter}{row_idx}']
            cell.font = regular_font
            cell.border = thin_border
            if fill.fill_type:
                cell.fill = fill
            if col_letter in ['A', 'B', 'D', 'E']:
                cell.alignment = Alignment(horizontal="center")
        row_idx += 1

    # Auto column width
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = col[0].column_letter
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()


def generate_attendance_pdf(centre_data, session_data, records_data):
    """Generate professional PDF for Attendance Report."""
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, rightMargin=36, leftMargin=36, topMargin=36, bottomMargin=36)
    elements = []
    styles = getSampleStyleSheet()

    # Custom styles
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontSize=18,
        leading=22,
        textColor=colors.HexColor('#1E3A8A'),
        spaceAfter=4
    )
    subtitle_style = ParagraphStyle(
        'DocSub',
        parent=styles['Normal'],
        fontSize=10,
        textColor=colors.HexColor('#475569'),
        spaceAfter=12
    )
    meta_style = ParagraphStyle(
        'Meta',
        parent=styles['Normal'],
        fontSize=9,
        textColor=colors.HexColor('#1E293B')
    )

    # Header
    elements.append(Paragraph("SKILL INDIA MONITORING & COMPLIANCE PLATFORM", subtitle_style))
    elements.append(Paragraph(f"Official Attendance & AI Cross-Check Verification", title_style))
    elements.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor('#1E3A8A'), spaceAfter=14))

    # Meta Table
    meta_data = [
        [
            Paragraph(f"<b>Training Centre:</b> {centre_data.get('centre_name')}", meta_style),
            Paragraph(f"<b>Centre ID:</b> {centre_data.get('centre_id')}", meta_style)
        ],
        [
            Paragraph(f"<b>Location:</b> {centre_data.get('location')}, {centre_data.get('state')}", meta_style),
            Paragraph(f"<b>Batch:</b> {session_data.get('batch', 'Solar-A')}", meta_style)
        ],
        [
            Paragraph(f"<b>Date / Session:</b> {session_data.get('date')} ({session_data.get('session_time')})", meta_style),
            Paragraph(f"<b>Verified By:</b> {session_data.get('submitted_by')}", meta_style)
        ]
    ]
    meta_table = Table(meta_data, colWidths=[3.5*inch, 3.5*inch])
    meta_table.setStyle(TableStyle([
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
    ]))
    elements.append(meta_table)
    elements.append(Spacer(1, 14))

    # AI vs Manual Comparison Callout Box
    manual_cnt = session_data.get('manual_present_count', 0)
    ai_cnt = session_data.get('ai_detected_count', 0)
    diff = session_data.get('difference_count', 0)
    pct = session_data.get('discrepancy_percentage', 0.0)

    summary_data = [
        ["Manual Recorded", "AI Camera Headcount", "Discrepancy Count", "Discrepancy %", "Status"],
        [str(manual_cnt), str(ai_cnt), f"{diff} Persons", f"{pct:.1f}%", session_data.get('discrepancy_status', 'NORMAL')]
    ]
    summary_table = Table(summary_data, colWidths=[1.4*inch, 1.4*inch, 1.4*inch, 1.4*inch, 1.4*inch])
    
    status_color = colors.HexColor('#DC2626') if session_data.get('discrepancy_status') == 'CRITICAL' else (colors.HexColor('#D97706') if session_data.get('discrepancy_status') == 'WARNING' else colors.HexColor('#16A34A'))
    summary_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1E3A8A')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 9),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('BACKGROUND', (0, 1), (-1, 1), colors.HexColor('#F8FAFC')),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
        ('TEXTCOLOR', (4, 1), (4, 1), status_color),
        ('FONTNAME', (0, 1), (-1, 1), 'Helvetica-Bold'),
        ('TOPPADDING', (0, 0), (-1, -1), 6),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
    ]))
    elements.append(summary_table)
    elements.append(Spacer(1, 16))

    # Trainees List Table
    elements.append(Paragraph("<b>Submitted Trainee Roster</b>", styles['Heading3']))
    elements.append(Spacer(1, 6))

    roster_rows = [["#", "Trainee ID", "Name", "Attendance", "Marked Time", "Remarks"]]
    for i, r in enumerate(records_data[:30], 1): # Top 30
        roster_rows.append([
            str(i),
            r.get('trainee_id', ''),
            r.get('name', ''),
            r.get('status', 'Present'),
            r.get('marked_time', '—'),
            r.get('remarks', '')[:25]
        ])

    roster_table = Table(roster_rows, colWidths=[0.4*inch, 1.2*inch, 2.0*inch, 1.1*inch, 1.0*inch, 1.3*inch])
    roster_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#3B82F6')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 8),
        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
        ('ALIGN', (0, 0), (0, -1), 'CENTER'),
        ('ALIGN', (3, 0), (4, -1), 'CENTER'),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F1F5F9')]),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E2E8F0')),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
    ]))
    elements.append(roster_table)

    # Privacy and Certification Note
    elements.append(Spacer(1, 14))
    elements.append(Paragraph(
        "<i>* Note on AI Person Detection: Headcount is computed via automated YOLO object detection without storing facial identities or biometric markers in compliance with Government Privacy Safeguards.</i>",
        ParagraphStyle('Foot', parent=styles['Normal'], fontSize=7.5, textColor=colors.HexColor('#64748B'))
    ))

    doc.build(elements)
    buf.seek(0)
    return buf.getvalue()


def generate_infrastructure_excel(centre_data, equipment_list, issues_list):
    """Generate professional Excel workbook for Infrastructure & Equipment Report."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Infrastructure Compliance"

    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    sub_fill = PatternFill(start_color="0284C7", end_color="0284C7", fill_type="solid")
    alt_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    bold_font = Font(name="Calibri", size=11, bold=True)
    regular_font = Font(name="Calibri", size=11)
    thin_border = Border(left=Side(style='thin', color='CBD5E1'), right=Side(style='thin', color='CBD5E1'),
                         top=Side(style='thin', color='CBD5E1'), bottom=Side(style='thin', color='CBD5E1'))

    ws['A1'] = "GOVERNMENT OF INDIA - SKILLING INFRASTRUCTURE VERIFICATION REPORT"
    ws['A1'].font = Font(name="Calibri", size=14, bold=True, color="1E3A8A")

    ws['A3'] = "Centre Name:"
    ws['B3'] = centre_data.get('centre_name', 'N/A')
    ws['A3'].font = bold_font
    ws['B3'].font = regular_font

    ws['A4'] = "Location:"
    ws['B4'] = f"{centre_data.get('location', '')}, {centre_data.get('state', '')}"
    ws['A4'].font = bold_font

    ws['D3'] = "Sanctioned Capacity:"
    ws['E3'] = f"{centre_data.get('capacity', 0)} Trainees"
    ws['D3'].font = bold_font

    ws['D4'] = "Report Date:"
    ws['E4'] = datetime.now().strftime("%Y-%m-%d")
    ws['D4'].font = bold_font

    # Equipment Table Header
    ws['A6'] = "Equipment ID"
    ws['B6'] = "Equipment Name"
    ws['C6'] = "Category"
    ws['D6'] = "Approved Qty"
    ws['E6'] = "AI Detected Qty"
    ws['F6'] = "Status"
    ws['G6'] = "Condition & Remarks"

    for col in ['A6', 'B6', 'C6', 'D6', 'E6', 'F6', 'G6']:
        ws[col].fill = header_fill
        ws[col].font = header_font
        ws[col].alignment = Alignment(horizontal="center" if col in ['A6', 'D6', 'E6', 'F6'] else "left")

    row_idx = 7
    for i, eq in enumerate(equipment_list):
        ws[f'A{row_idx}'] = eq.get('equipment_id', '')
        ws[f'B{row_idx}'] = eq.get('name', '')
        ws[f'C{row_idx}'] = eq.get('category', '')
        ws[f'D{row_idx}'] = eq.get('approved_quantity', 0)
        ws[f'E{row_idx}'] = eq.get('current_detected_quantity', 0)
        ws[f'F{row_idx}'] = eq.get('status', 'Functional')
        ws[f'G{row_idx}'] = f"{eq.get('condition_desc', '')} {eq.get('remarks', '')}".strip()

        fill = alt_fill if i % 2 == 0 else PatternFill(fill_type=None)
        for col_letter in ['A', 'B', 'C', 'D', 'E', 'F', 'G']:
            cell = ws[f'{col_letter}{row_idx}']
            cell.font = regular_font
            cell.border = thin_border
            if fill.fill_type:
                cell.fill = fill
            if col_letter in ['A', 'D', 'E', 'F']:
                cell.alignment = Alignment(horizontal="center")
        row_idx += 1

    # Open Issues Table
    row_idx += 2
    ws[f'A{row_idx}'] = "REPORTED DEFICIENCIES & COMMENTS"
    ws[f'A{row_idx}'].font = Font(name="Calibri", size=13, bold=True, color="1E3A8A")
    row_idx += 1

    ws[f'A{row_idx}'] = "Issue ID"
    ws[f'B{row_idx}'] = "Category"
    ws[f'C{row_idx}'] = "Severity"
    ws[f'D{row_idx}'] = "Description"
    ws[f'E{row_idx}'] = "Submitted By"
    ws[f'F{row_idx}'] = "Status"
    ws[f'G{row_idx}'] = "Government Action"

    for col_letter in ['A', 'B', 'C', 'D', 'E', 'F', 'G']:
        ws[f'{col_letter}{row_idx}'].fill = sub_fill
        ws[f'{col_letter}{row_idx}'].font = header_font
        ws[f'{col_letter}{row_idx}'].alignment = Alignment(horizontal="center" if col_letter in ['A', 'C', 'F'] else "left")

    row_idx += 1
    for i, iss in enumerate(issues_list):
        ws[f'A{row_idx}'] = iss.get('issue_id', '')
        ws[f'B{row_idx}'] = iss.get('category', '')
        ws[f'C{row_idx}'] = iss.get('severity', '')
        ws[f'D{row_idx}'] = iss.get('description', '')
        ws[f'E{row_idx}'] = iss.get('submitted_by', '')
        ws[f'F{row_idx}'] = iss.get('status', '')
        ws[f'G{row_idx}'] = iss.get('government_notes', 'Pending review')

        for col_letter in ['A', 'B', 'C', 'D', 'E', 'F', 'G']:
            cell = ws[f'{col_letter}{row_idx}']
            cell.font = regular_font
            cell.border = thin_border
        row_idx += 1

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = col[0].column_letter
        ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()


def generate_compliance_pdf(centre_data, session_data, equipment_list, alerts_list, config_data):
    """Generate comprehensive Centre Compliance & Monitoring Report."""
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, rightMargin=36, leftMargin=36, topMargin=36, bottomMargin=36)
    elements = []
    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontSize=18,
        leading=22,
        textColor=colors.HexColor('#1E3A8A'),
        spaceAfter=4
    )
    subtitle_style = ParagraphStyle(
        'DocSub',
        parent=styles['Normal'],
        fontSize=10,
        textColor=colors.HexColor('#475569'),
        spaceAfter=12
    )
    meta_style = ParagraphStyle('Meta', parent=styles['Normal'], fontSize=9, textColor=colors.HexColor('#1E293B'))

    elements.append(Paragraph("SKILL DEVELOPMENT COMPLIANCE AUDIT COMMISSION", subtitle_style))
    elements.append(Paragraph(f"Comprehensive Centre Compliance & AI Verification Dossier", title_style))
    elements.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor('#1E3A8A'), spaceAfter=14))

    score = centre_data.get('compliance_score', 80.0)
    rating = "EXCELLENT" if score >= 90 else ("GOOD" if score >= 75 else ("NEEDS ATTENTION" if score >= 50 else "CRITICAL"))
    rating_color = colors.HexColor('#16A34A') if score >= 75 else (colors.HexColor('#D97706') if score >= 50 else colors.HexColor('#DC2626'))

    # Overview box
    summary_data = [
        [
            Paragraph(f"<b>Centre Name:</b> {centre_data.get('centre_name')}<br/><b>Location:</b> {centre_data.get('location')}, {centre_data.get('state')}<br/><b>Sanctioned Capacity:</b> {centre_data.get('capacity')} Trainees", meta_style),
            Paragraph(f"<b>Overall Compliance Score:</b> <font size='14' color='{rating_color.hexval()}'><b>{score:.1f} / 100</b></font><br/><b>Compliance Tier:</b> <b>{rating}</b><br/><b>Formula:</b> Attendance {int(config_data.get('attendance_weight', 0.4)*100)}% + Infra {int(config_data.get('infrastructure_weight', 0.4)*100)}% + Issues {int(config_data.get('issue_weight', 0.2)*100)}%", meta_style)
        ]
    ]
    summary_table = Table(summary_data, colWidths=[3.5*inch, 3.5*inch])
    summary_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#F8FAFC')),
        ('BOX', (0, 0), (-1, -1), 1, colors.HexColor('#CBD5E1')),
        ('PADDING', (0, 0), (-1, -1), 10),
    ]))
    elements.append(summary_table)
    elements.append(Spacer(1, 14))

    # Active Alerts & Discrepancies
    elements.append(Paragraph("<b>Flagged Discrepancies & Audit Alerts</b>", styles['Heading3']))
    elements.append(Spacer(1, 6))

    alert_rows = [["Alert ID", "Type", "Severity", "Description", "Status"]]
    for al in alerts_list:
        alert_rows.append([
            al.get('alert_id', ''),
            al.get('alert_type', ''),
            al.get('severity', ''),
            Paragraph(al.get('description', '')[:90] + "...", ParagraphStyle('ALDesc', parent=styles['Normal'], fontSize=8)),
            al.get('status', '')
        ])

    alert_table = Table(alert_rows, colWidths=[1.1*inch, 1.5*inch, 0.9*inch, 2.5*inch, 1.0*inch])
    alert_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#B91C1C')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 8),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#FEF2F2')]),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#FCA5A5')),
        ('TOPPADDING', (0, 0), (-1, -1), 5),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
    ]))
    elements.append(alert_table)
    elements.append(Spacer(1, 14))

    # Infrastructure status summary
    elements.append(Paragraph("<b>Infrastructure & Equipment Audit Status</b>", styles['Heading3']))
    elements.append(Spacer(1, 6))

    eq_rows = [["ID", "Equipment Name", "Approved", "Detected", "Discrepancy", "Status"]]
    for eq in equipment_list:
        diff_eq = eq.get('approved_quantity', 0) - eq.get('current_detected_quantity', 0)
        eq_rows.append([
            eq.get('equipment_id', ''),
            eq.get('name', '')[:25],
            str(eq.get('approved_quantity', 0)),
            str(eq.get('current_detected_quantity', 0)),
            f"{diff_eq} missing" if diff_eq > 0 else "OK",
            eq.get('status', 'Functional')
        ])

    eq_table = Table(eq_rows, colWidths=[1.0*inch, 2.4*inch, 0.9*inch, 0.9*inch, 1.0*inch, 0.8*inch])
    eq_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1E3A8A')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 8),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F8FAFC')]),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
    ]))
    elements.append(eq_table)

    elements.append(Spacer(1, 14))
    elements.append(Paragraph(
        "<i>Prototype Compliance Indicator: This report is generated dynamically by the AI Monitoring System prototype. Final legal validation requires on-site inspection and authorized physical officer sign-off.</i>",
        ParagraphStyle('Foot', parent=styles['Normal'], fontSize=7.5, textColor=colors.HexColor('#64748B'))
    ))

    doc.build(elements)
    buf.seek(0)
    return buf.getvalue()
